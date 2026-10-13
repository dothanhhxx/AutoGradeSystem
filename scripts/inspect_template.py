"""Inspect a PPTX template: slide size, layouts, placeholders, and existing slides."""
import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1] if len(sys.argv) > 1 else "ICOCO2026-presentation-template.pptx"
prs = Presentation(path)
print(f"Slide size: {Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in")
print(f"Slides: {len(prs.slides)}")

print("\n=== LAYOUTS ===")
for mi, master in enumerate(prs.slide_masters):
    print(f"Master {mi}: {len(master.slide_layouts)} layouts")
    for li, layout in enumerate(master.slide_layouts):
        print(f"  [{li}] {layout.name}")
        for ph in layout.placeholders:
            pf = ph.placeholder_format
            print(f"       ph idx={pf.idx} type={pf.type} name={ph.name!r} "
                  f"pos=({Emu(ph.left).inches:.2f},{Emu(ph.top).inches:.2f}) "
                  f"size=({Emu(ph.width).inches:.2f},{Emu(ph.height).inches:.2f})")

print("\n=== SLIDES ===")
for si, slide in enumerate(prs.slides):
    print(f"Slide {si + 1}: layout={slide.slide_layout.name!r}")
    for sh in slide.shapes:
        txt = sh.text_frame.text.replace("\n", " | ")[:100] if sh.has_text_frame else ""
        ph = f"ph={sh.placeholder_format.idx}" if sh.is_placeholder else "-"
        print(f"    {sh.shape_type} {sh.name!r} {ph} "
              f"pos=({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) "
              f"size=({Emu(sh.width).inches:.2f},{Emu(sh.height).inches:.2f}) {txt!r}")
